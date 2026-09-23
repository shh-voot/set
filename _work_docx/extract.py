import sys, docx
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

path = r"D:\claude\华为杯\D题\_work_docx\problem.docx"
doc = Document(path)

def iter_block_items(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

out = []
ti = 0
for block in iter_block_items(doc):
    if isinstance(block, Paragraph):
        t = block.text.strip()
        style = block.style.name if block.style is not None else ''
        if t:
            out.append(f"[P|{style}] {t}")
    else:
        ti += 1
        out.append(f"===== TABLE {ti} =====")
        for r in block.rows:
            cells = [c.text.strip().replace('\n', ' / ') for c in r.cells]
            out.append(" | ".join(cells))
        out.append(f"===== END TABLE {ti} =====")

text = "\n".join(out)
with open(r"D:\claude\华为杯\D题\_work_docx\problem_text.txt", "w", encoding="utf-8") as f:
    f.write(text)

print("paragraphs:", len(doc.paragraphs))
print("tables:", len(doc.tables))
print("chars:", len(text))
# inline shapes (images)
print("inline_shapes:", len(doc.inline_shapes))
