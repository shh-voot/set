import openpyxl, os, glob

base = r"D:\claude\华为杯\D题\数据\无人机应急物资运输基础数据"
out = []
for p in sorted(glob.glob(os.path.join(base, "*.xlsx"))):
    out.append("=" * 100)
    out.append("FILE: " + os.path.basename(p))
    out.append("=" * 100)
    wb = openpyxl.load_workbook(p, data_only=True)
    for ws in wb.worksheets:
        out.append(f"----- SHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column} -----")
        n = 0
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if all(v == "" for v in vals):
                continue
            out.append(" | ".join(vals))
            n += 1
            if n > 300:
                out.append("... (truncated)")
                break
        # merged cells
        if ws.merged_cells.ranges:
            out.append("MERGED: " + ", ".join(str(r) for r in ws.merged_cells.ranges))
    out.append("")

with open(r"D:\claude\华为杯\D题\_work_docx\xlsx_dump.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("done", len(out))
