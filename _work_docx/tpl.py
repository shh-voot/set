import openpyxl
p = r"D:\claude\华为杯\D题\结果提交模板.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
for ws in wb.worksheets:
    print("="*90)
    print("SHEET:", ws.title, ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
    for row in ws.iter_rows(values_only=True):
        vals = ["" if v is None else str(v) for v in row]
        if all(v=="" for v in vals): continue
        print(" | ".join(vals))
    if ws.merged_cells.ranges:
        print("MERGED:", ", ".join(str(r) for r in ws.merged_cells.ranges))
