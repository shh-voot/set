"""Package existing audit evidence and verify original file integrity."""
from pathlib import Path
import hashlib
import json
import subprocess
import sys
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
E = HERE / 'evidence'
report = (HERE / '工作质量与公式溯源审查报告.md').read_text(encoding='utf-8')
formulas = []
for line in report.splitlines():
    if line.startswith('| F'):
        parts = [s.strip().replace('**', '').replace('`', '') for s in line.strip('|').split('|')]
        formulas.append(dict(zip(['编号', '公式或规则', '来源定位', '审查结论'], parts)))
summary = json.loads((E/'audit_summary.json').read_text(encoding='utf-8'))
targeted = json.loads((E/'targeted_checks.json').read_text(encoding='utf-8'))
manifest = json.loads((E/'input_manifest.json').read_text(encoding='utf-8'))
integrity = []
for entry in manifest:
    p = ROOT / entry['path']
    actual = hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None
    integrity.append(dict(path=entry['path'], original_sha256=entry['sha256'], current_sha256=actual,
                          unchanged=actual == entry['sha256']))
check = dict(checked_files=len(integrity), changed_files=[r['path'] for r in integrity if not r['unchanged']],
             files=integrity)
(E/'integrity_check.json').write_text(json.dumps(check, ensure_ascii=False, indent=2), encoding='utf-8')
assert not check['changed_files'], check['changed_files']

notes = [
    ('审查版本', '2026-09-25，ea5e0fb；请先阅读同目录中文主报告。'),
    ('结论', '不能作为已通过全部题目约束的终版；通信、中继续航、问题四时限等已有失败证据。'),
    ('原数据保护', f'{len(integrity)} 个原有受跟踪文件 SHA256 与审计开始时一致。'),
    ('original / published', '保留交付表的字段和数值；其中 feasible、on_time 为团队原字段，不代表本审计认可。'),
    ('code_', '调用现有公共函数重算，检验结果能否复现；不证明公式正确。'),
    ('exponent_only_', '只将 L(q) 改为题给 3/2 次幂，保留其他计算口径与原组批。'),
    ('reference_', '明确假设 Ehor=C*d/L；Eup=mgh/eta；逐像元取高。用于敏感性对照，不是唯一官方修正答案。'),
    ('reference_duration', '三阶段题给时间式，逐像元几何，保留原开始时刻；能耗假设不影响该时长。'),
    ('exact_dem / exact_cruise', '直线按栅格边界切分，遍历每个相交像元；恰好触角且交长为零的格子不计。'),
    ('dem_underestimate_m', '正确像元最大值减代码值，正数为代码低估。'),
    ('load_ok / volume_ok', '按附件逐箱重算后检查，不读取原表自报合规结论。'),
    ('hard_late', '医疗期望时间与首批截止的硬约束；其余期望时间是软及时性指标。'),
    ('通信余量', '单位 dB，相对包含衰落裕量的门限；负数为链路不满足。中继列为接入/回传较差者。'),
    ('通信验证边界', '真实投送点反例足以否定连续通信；有限路径采样未声称证明剩余时刻全部通过。'),
    ('资源独立统计', '固定当前 Q3 时刻与交付分区，按每组半开占用区间峰值求和；不代表最终最优配置。'),
    ('时间与能量', '审计工作表时间采用 min，除表头另有声明；能量 kWh，高度 m。官方提交模板另要求 s。'),
    ('复现差异', 'Q4 两个独立资源工作簿共 4 个 sheet 不一致；详见 targeted_checks.json。'),
]
csvs = {
    'Q1逐架次': 'q1_mission_check.csv', 'Q2逐架次': 'q2_mission_check.csv',
    'Q3逐架次': 'q3_mission_check.csv', 'Q2逐箱': 'q2_cargo_check.csv',
    'Q3逐箱': 'q3_cargo_check.csv', 'Q3路径通信': 'q3_communication_check.csv',
    'Q3投送点交叉验证': 'q3_endpoint_robustness.csv', 'Q3断连采样点': 'q3_failed_points.csv',
    'Q4两组硬迟到': 'q4_2_hard_late.csv', 'Q4三组硬迟到': 'q4_3_hard_late.csv',
    'Q4继承Q3资源': 'q4_from_actual_q3_times.csv',
}
tables = {'阅读说明': pd.DataFrame(notes, columns=['项目', '说明']), '公式溯源': pd.DataFrame(formulas),
          '中继续航': pd.DataFrame(summary['relay_endurance']),
          '余量敏感性反例': pd.DataFrame(targeted['reserve_sensitivity'])}
tables.update({name: pd.read_csv(E/file) for name, file in csvs.items()})
for q in ('q2','q3'):
    for kind in ('uav', 'battery'):
        tables[f'{q.upper()}对照{kind}冲突'] = pd.DataFrame(summary[q][f'reference_{kind}_conflicts'])
tables['复现状态'] = pd.DataFrame(json.loads((E/'reproduction_status.json').read_text(encoding='utf-8')))
tables['原文件哈希'] = pd.DataFrame(integrity)
with pd.ExcelWriter(HERE/'审查明细.xlsx', engine='openpyxl') as writer:
    for sheet, table in tables.items():
        table.to_excel(writer, sheet_name=sheet, index=False)
        ws = writer.sheets[sheet]
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.zoomScale = 85
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='234862')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[1].height = 40
        for cells in ws.columns:
            letter = cells[0].column_letter
            longest = max(len(str(c.value or '')) for c in cells)
            ws.column_dimensions[letter].width = min(60, max(15, longest + 2))
            for cell in cells[1:]:
                cell.alignment = Alignment(vertical='top', wrap_text=sheet in ('阅读说明','公式溯源'))
                if isinstance(cell.value, float):
                    cell.number_format = '0.000000'
        if sheet in ('阅读说明','公式溯源'):
            for row in range(2, ws.max_row+1):
                ws.row_dimensions[row].height = 58

freeze = subprocess.check_output([sys.executable, '-m', 'pip', 'freeze'], text=True)
(HERE/'requirements-audit.txt').write_text(freeze, encoding='utf-8')
print(json.dumps(dict(workbook_sheets=len(tables), formulas=len(formulas),
                     checked_original_files=len(integrity), changed_original_files=check['changed_files']), ensure_ascii=False))
