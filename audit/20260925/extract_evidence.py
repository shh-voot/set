"""Extract primary files for the audit; never modifies the delivered work."""
from pathlib import Path
import hashlib
import json
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / 'evidence'
OUT.mkdir(parents=True, exist_ok=True)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

def mathtext(e):
    tag = e.tag.split('}')[-1]
    def child(n):
        c = e.find(M+n)
        return mathtext(c) if c is not None else ''
    if tag.endswith('Pr'):
        return ''
    if tag == 't':
        return e.text or ''
    if tag == 'sSub':
        return child('e') + '_{' + child('sub') + '}'
    if tag == 'sSup':
        return child('e') + '^{' + child('sup') + '}'
    if tag == 'sSubSup':
        return child('e') + '_{' + child('sub') + '}^{' + child('sup') + '}'
    if tag == 'f':
        return '(' + child('num') + ')/(' + child('den') + ')'
    if tag == 'd':
        return '(' + child('e') + ')'
    if tag == 'nary':
        return 'SUM_{'+child('sub')+'}^{'+child('sup')+'} '+child('e')
    if tag == 'mr':
        return ' [' + ' | '.join(mathtext(c) for c in e) + '] '
    return ''.join(mathtext(c) for c in e)

doc = ROOT / '山区洪涝灾害下无人机运输与通信协同优化.docx'
with zipfile.ZipFile(doc) as z:
    xml = z.read('word/document.xml')
    (OUT/'problem_document.xml').write_bytes(xml)
    tree = ET.fromstring(xml)
    paras = []
    for i, p in enumerate(tree.iter(W+'p'), 1):
        def walk(e):
            if e.tag == M+'oMath':
                return ' $' + mathtext(e) + '$ '
            if e.tag == W+'t':
                return e.text or ''
            return ''.join(walk(c) for c in e)
        s = walk(p).strip()
        if s:
            paras.append(f'P{i:03d}: {s}')
    (OUT/'problem_with_formulas.txt').write_text('\n'.join(paras), encoding='utf-8')
    for name in z.namelist():
        if name.startswith('word/media/'):
            dest = OUT/'docx_media'/Path(name).name
            dest.parent.mkdir(exist_ok=True)
            dest.write_bytes(z.read(name))

import openpyxl
books = list((ROOT/'数据/无人机应急物资运输基础数据').glob('*.xlsx'))
books += [ROOT/'结果提交模板.xlsx']
books += list((ROOT/'解题-gpt/结果').glob('*.xlsx'))
summaries = []
for path in books:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    sheet_meta = []
    for ws in wb:
        lines.append(f'SHEET {ws.title} ({ws.max_row} x {ws.max_column})')
        sheet_meta.append({'name': ws.title, 'rows': ws.max_row, 'cols': ws.max_column})
        for row in ws:
            cells = [f'{c.coordinate}={c.value}' for c in row if c.value is not None]
            if cells:
                lines.append(' | '.join(cells))
    folder = OUT / ('official_xlsx' if path in books[:6] else 'results_xlsx')
    folder.mkdir(exist_ok=True)
    (folder/(path.stem+'.txt')).write_text('\n'.join(lines), encoding='utf-8')
    summaries.append({'file': str(path.relative_to(ROOT)), 'sheets': sheet_meta})
(OUT/'workbook_inventory.json').write_text(json.dumps(summaries, ensure_ascii=False, indent=2), encoding='utf-8')

try:
    import fitz
except ImportError:
    fitz = None
pdfs = []
if fitz:
    for path in sorted((ROOT/'文献库').glob('*.pdf')):
        info = {'file': path.name, 'bytes': path.stat().st_size}
        try:
            d = fitz.open(path)
            info.update(pages=len(d), metadata=d.metadata)
            texts = [f'--- PDF PAGE {i+1} ---\n'+p.get_text() for i,p in enumerate(d)]
            dest = OUT/'papers'/f'{path.stem}.txt'
            dest.parent.mkdir(exist_ok=True)
            dest.write_text('\n'.join(texts), encoding='utf-8')
            info['first_page'] = texts[0][:6000]
        except Exception as e:
            info['error'] = str(e)
        pdfs.append(info)
    (OUT/'pdf_inventory.json').write_text(json.dumps(pdfs, ensure_ascii=False, indent=2), encoding='utf-8')
    gp = ROOT/'数据/镇龙乡地理空间数据/镇龙乡地理空间数据说明.pdf'
    with fitz.open(gp) as d:
        (OUT/'geo_description.txt').write_text('\n'.join(p.get_text() for p in d), encoding='utf-8')

tracked = subprocess.check_output(['git','ls-files','-z'], cwd=ROOT).decode('utf-8').split('\0')
manifest = []
for f in tracked:
    if not f:
        continue
    p = ROOT/f
    manifest.append({'path': f, 'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()})
(OUT/'input_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'tracked_files':len(manifest),'workbooks':len(books),'pdfs':len(pdfs),'evidence_dir':str(OUT)},ensure_ascii=False))
