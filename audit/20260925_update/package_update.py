from pathlib import Path
import json,hashlib,subprocess,sys
import pandas as pd
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent;E=HERE/'evidence'
text=(HERE/'最新提交补充评估报告.md').read_text(encoding='utf8')
def markdown_table(header):
    lines=text.splitlines();start=next(i for i,x in enumerate(lines) if x.startswith(header));rows=[]
    keys=[x.strip() for x in lines[start].strip('|').split('|')]
    for line in lines[start+2:]:
        if not line.startswith('|'):break
        rows.append([x.strip().replace('**','').replace('`','') for x in line.strip('|').split('|')])
    return pd.DataFrame(rows,columns=keys)
tables={'上轮问题状态':markdown_table('| 上轮问题 |'),'新增公式核对':markdown_table('| 新公式/论述 |')}
notes=[('审查版本','e94e1cc；完成日期2026-09-26。上一轮ea5e0fb记录原样保留。'),
       ('结论','仍需实质性修订；新增Q3实际超载，公共模型和通信/中继续航问题仍在。'),
       ('字段来源','feasible、on_time等无审计前缀字段来自原交付，不代表本轮审计认可。'),
       ('code_*','用当前公共函数检验复算一致，不证明该物理函数符合题目。'),
       ('reference_*','工程假设对照Ehor=C*d/L、Eup=mgh/eta及逐像元几何，非唯一官方修正值。'),
       ('Q3定义域','超载行超出题给L(q)定义域。reference能耗的外推无合规意义，不能用于修正答案。'),
       ('硬截止','医疗期望截止、首批保障截止；普通物资期望时限另为软指标。'),
       ('通信','实际投送点反例+交叉验证；有限采样不证明全过程连续可用。'),
       ('合成CP反例','1.09min等案例是控制变量测试，不是官方货箱实际结果。'),
       ('ALNS复现','主要指标复现；具体37/43行任务顺序不同，不能声称逐表完全一致。'),
       ('Q4对照资源','固定有错误的Q3时刻和交付分区作继承诊断，不是新合规配置。')]
tables={'阅读说明':pd.DataFrame(notes,columns=['项目','说明']),**tables}
mapping={'Q1逐架次':'q1_mission_check.csv','Q2逐架次':'q2_mission_check.csv','Q3逐架次':'q3_mission_check.csv',
         'Q2逐箱':'q2_cargo_check.csv','Q3逐箱':'q3_cargo_check.csv',
         '投送点通信交叉验证':'q3_endpoint_robustness.csv','路径通信反例':'q3_failed_points.csv',
         'Q4两组硬迟到':'q4_2_hard_late.csv','Q4三组硬迟到':'q4_3_hard_late.csv',
         'Q4继承Q3对照':'q4_from_actual_q3_times.csv','新增候选逐架次':'new_candidate_checks.csv',
         'Q1容量下界':'q1_capacity_lower_bound.csv'}
tables.update({key:pd.read_csv(E/file) for key,file in mapping.items()})
q3=tables['Q3逐架次'];tables['Q3实际超载']=q3.loc[~q3.load_ok|~q3.volume_ok]
summary=json.loads((E/'audit_summary.json').read_text(encoding='utf8'))
tables['中继单组件续航']=pd.DataFrame(summary['relay_endurance'])
target=json.loads((E/'targeted_checks.json').read_text(encoding='utf8'))
tables['主链复现差异']=pd.DataFrame([{k:v for k,v in r.items() if k not in ('reproduced','delivered')} for r in target['reproduction_differences']])
tables['余量敏感性']=pd.DataFrame(target['reserve_sensitivity'])
add=json.loads((E/'additions_summary.json').read_text(encoding='utf8'))
tables['新增能耗辅助函数']=pd.DataFrame(add['claude_energy'])
tables['CP截断迟到反例']=pd.DataFrame(add['cp_truncation_lateness']['missions'])
tables['CP截断重叠反例']=pd.DataFrame(add['cp_truncation_overlap']['missions'])
tables['运行记录']=pd.DataFrame(json.loads((E/'reproduction_status.json').read_text(encoding='utf8'))+
                               json.loads((E/'extra_run_status.json').read_text(encoding='utf8')))
pre=json.loads((HERE/'pull_preflight.json').read_text(encoding='utf8'))
post=json.loads((E/'post_pull_manifest.json').read_text(encoding='utf8'))
def verify(entries):
    rows=[]
    for item in entries:
        p=ROOT/item['path'];actual=hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None
        rows.append(dict(path=item['path'],sha256=item['sha256'],current_sha256=actual,unchanged=actual==item['sha256']))
    return rows
oldcheck=verify(pre['prior_audit_manifest']);newcheck=verify(post)
assert all(r['unchanged'] for r in oldcheck+newcheck)
integrity=dict(prior_audit_files=len(oldcheck),prior_audit_changed=[],new_tracked_files=len(newcheck),new_tracked_changed=[],
               before_head=pre['before_head'],after_head=pre['remote_head'],path_collisions=pre['audit_path_collisions'],
               prior_audit=oldcheck,tracked=newcheck)
(E/'integrity_check.json').write_text(json.dumps(integrity,ensure_ascii=False,indent=2),encoding='utf8')
tables['原审计保留核对']=pd.DataFrame(oldcheck);tables['新版本文件保护']=pd.DataFrame(newcheck)
with pd.ExcelWriter(HERE/'补充评估明细.xlsx',engine='openpyxl') as writer:
    for sheet,table in tables.items():
        table.to_excel(writer,sheet_name=sheet,index=False);ws=writer.sheets[sheet]
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.zoomScale=85
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234862');cell.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=40
        for cells in ws.columns:
            ws.column_dimensions[cells[0].column_letter].width=min(62,max(15,max(len(str(c.value or '')) for c in cells)+2))
            for cell in cells[1:]:
                cell.alignment=Alignment(vertical='top',wrap_text=sheet in ('阅读说明','上轮问题状态','新增公式核对'))
                if isinstance(cell.value,float):cell.number_format='0.000000'
        if sheet in ('阅读说明','上轮问题状态','新增公式核对'):
            for i in range(2,ws.max_row+1):ws.row_dimensions[i].height=58
(HERE/'requirements-audit.txt').write_text(subprocess.check_output([sys.executable,'-m','pip','freeze'],text=True),encoding='utf8')
print(json.dumps(dict(workbook_sheets=len(tables),prior_audit_preserved=len(oldcheck),tracked_preserved=len(newcheck)),ensure_ascii=False))
